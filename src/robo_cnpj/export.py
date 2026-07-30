"""Exportação para Excel — o contador ainda vive na planilha."""

from __future__ import annotations

from pathlib import Path

from rich.console import Console

from . import db
from .config import Config

console = Console()

COLUNAS = [
    ("cnpj_formatado", "CNPJ", 20),
    ("razao_social", "Razão social", 45),
    ("nome_fantasia", "Nome fantasia", 30),
    ("municipio", "Município", 22),
    ("regime_provavel", "Regime provável", 40),
    ("regime_confianca", "Confiança", 12),
    ("porte", "Porte", 24),
    ("capital_social", "Capital social", 16),
    ("situacao", "Situação", 12),
    ("cnae_descricao", "Atividade principal", 50),
    ("score", "Score", 8),
    ("apontamentos_texto", "Apontamentos", 70),
    ("cnd_federal_prognostico", "Prognóstico CND federal", 45),
    ("dau_qtd_inscricoes", "Inscrições DAU", 14),
    ("dau_valor_total", "Valor DAU (R$)", 18),
    ("cdt_estadual_mg_prognostico", "Prognóstico CDT estadual (MG)", 45),
    ("dam_qtd_inscricoes", "Inscrições dívida MG", 16),
    ("dam_valor_total", "Valor dívida MG (R$)", 18),
    ("telefone", "Telefone", 18),
    ("email", "E-mail", 32),
]


def executar(cfg: Config, destino: Path | None = None, limite: int | None = None,
             minimo_apontamentos: int = 0) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    saida = destino or (cfg.caminhos.exports / f"leads_{cfg.uf.lower()}.xlsx")
    saida.parent.mkdir(parents=True, exist_ok=True)

    campos = ", ".join(c for c, _, _ in COLUNAS)
    sql = f"""
        SELECT {campos} FROM leads
        WHERE qtd_apontamentos >= {minimo_apontamentos}
        ORDER BY score DESC, dau_valor_total DESC, razao_social
    """
    if limite:
        sql += f" LIMIT {limite}"

    with db.conectar(cfg, somente_leitura=True) as con:
        linhas = con.execute(sql).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Leads {cfg.uf}"

    cabecalho_fill = PatternFill("solid", fgColor="1F3864")
    cabecalho_font = Font(color="FFFFFF", bold=True)

    for i, (_, titulo, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=1, column=i, value=titulo)
        celula.fill = cabecalho_fill
        celula.font = cabecalho_font
        celula.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largura

    for linha in linhas:
        ws.append(list(linha))

    # Formato de moeda nas colunas de valor.
    for nome_campo, indice in ((c, i) for i, (c, _, _) in enumerate(COLUNAS, start=1)):
        if nome_campo in ("capital_social", "dau_valor_total", "dam_valor_total"):
            letra = get_column_letter(indice)
            for celula in ws[letra][1:]:
                celula.number_format = '#,##0.00'

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(saida)

    console.print(f"[green]✓[/] {len(linhas):,} leads → {saida}".replace(",", "."))
    return saida
