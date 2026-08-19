"""Teste de ponta a ponta do pipeline com a amostra sintética.

Roda sem internet. Executa: carga dos CSVs → transformação → verificação das
regras de negócio → export da planilha.

    python tests/test_pipeline.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ / "tests"))

import amostra  # noqa: E402
from robo_cnpj import config, db, export, schemas, transform  # noqa: E402
from robo_cnpj.ingest import mg_divida, pgfn, transparencia  # noqa: E402

falhas: list[str] = []


def conferir(condicao: bool, descricao: str, detalhe: str = "") -> None:
    if condicao:
        print(f"  [ok]    {descricao}")
    else:
        print(f"  [FALHA] {descricao}" + (f" — {detalhe}" if detalhe else ""))
        falhas.append(descricao)


def main() -> int:
    trabalho = Path(tempfile.mkdtemp(prefix="robo_cnpj_teste_"))
    print(f"Diretório de trabalho: {trabalho}\n")

    try:
        cfg = config.carregar()
        # Redireciona banco e exports para o temporário, preservando o banco real.
        object.__setattr__(cfg.caminhos, "banco", trabalho / "teste.duckdb")
        object.__setattr__(cfg.caminhos, "exports", trabalho / "exports")

        print("1. Gerando amostra sintética")
        arquivos = amostra.gerar(trabalho / "csv")
        print(f"   {len(arquivos)} arquivos\n")

        print("2. Carregando na base (mesmo caminho de código da produção)")
        with db.conectar(cfg) as con:
            for tabela in schemas.TODAS_TABELAS:
                db.criar_tabela(con, tabela)
                filtro = cfg.uf if "uf" in tabela.nomes_colunas else None
                n = db.carregar_csv(con, tabela, arquivos[tabela.nome], filtro_uf=filtro)
                print(f"   {tabela.nome}: {n} linhas")

            pgfn._criar_tabela(con)
            n = pgfn._carregar(con, arquivos["dau"], "Dívida Ativa - Geral", cfg.uf)
            print(f"   divida_ativa_uniao: {n} linhas")

            transparencia._criar_tabela(con)
            n = transparencia._carregar(con, arquivos["sancoes"], "CEIS")
            print(f"   sancoes: {n} linhas")

            mg_divida._criar_tabela(con)
            corte = mg_divida.corte_serial_para(cfg.mg_divida_janela_dias)
            n = mg_divida._inserir(con, amostra.divida_mg_linhas(), corte)
            print(f"   divida_ativa_mg: {n} linhas (de 3 na fixture — 1 fica fora da janela)")

        print("\n3. Transformando (filtro de apontamento OFF, para checar regra por linha)")
        # A regra de negócio (regime, score, prognóstico) precisa ser validada em
        # cada uma das 5 empresas do universo, incluindo o banco — que não tem
        # nenhum apontamento e por isso desaparece de `leads` quando o filtro
        # padrão está ligado. Roda uma vez sem o filtro para essas checagens, e
        # depois liga o filtro (produção) para checar o corte em si.
        object.__setattr__(cfg, "somente_com_apontamento", False)
        transform.executar(cfg)

        print("\n4. Verificando as regras de negócio")
        with db.conectar(cfg, somente_leitura=True) as con:
            def uma(sql: str, *params):
                return con.execute(sql, list(params)).fetchone()

            def existe(cnpj_basico: str) -> bool:
                return uma("SELECT count(*) FROM leads WHERE cnpj_basico = ?",
                           cnpj_basico)[0] > 0

            total = uma("SELECT count(*) FROM leads")[0]

            # Universo: 9 empresas na amostra. Devem sair:
            #   10000002 Simples ativo · 10000005 MEI ativo
            #   10000008 natureza 3999 (associação) · 10000009 UF=SP
            # Sobram 5: 01, 03, 04, 06, 07. A filial de 10000001 não entra.
            conferir(total == 5, f"universo tem 5 leads (obtido: {total})")

            conferir(not existe("10000002"), "Simples ativo é excluído")
            conferir(not existe("10000005"), "MEI ativo é excluído")
            conferir(not existe("10000008"), "associação (natureza 3xxx) é excluída")
            conferir(not existe("10000009"), "empresa de outra UF é excluída")

            filiais = uma("SELECT count(*) FROM leads WHERE cnpj = '10000001000236'")[0]
            conferir(filiais == 0, "filial não entra no universo")

            # Regime
            regime_banco = uma("SELECT regime_provavel, regime_confianca FROM leads "
                               "WHERE cnpj_basico = '10000003'")
            conferir(regime_banco[0].startswith("Lucro Real (obrigatório"),
                     "CNAE de banco → Lucro Real obrigatório", str(regime_banco))
            conferir(regime_banco[1] == "Alta", "confiança Alta no caso obrigatório")

            regime_metal = uma("SELECT regime_provavel FROM leads "
                               "WHERE cnpj_basico = '10000001'")[0]
            conferir(regime_metal == "Lucro Real (provável)",
                     "porte Demais + capital 5M → Lucro Real provável", regime_metal)

            regime_transp = uma("SELECT regime_provavel FROM leads "
                                "WHERE cnpj_basico = '10000004'")[0]
            conferir(regime_transp == "Lucro Presumido (provável)",
                     "porte Demais + capital 300k → Lucro Presumido", regime_transp)

            regime_cons = uma("SELECT regime_provavel FROM leads "
                              "WHERE cnpj_basico = '10000006'")[0]
            conferir("ME/EPP fora do Simples" in regime_cons,
                     "EPP ex-Simples → Presumido ME/EPP", regime_cons)

            # Dívida ativa: soma das duas inscrições da metalúrgica.
            dau = uma("SELECT dau_qtd_inscricoes, dau_valor_total, dau_ajuizado "
                      "FROM leads WHERE cnpj_basico = '10000001'")
            conferir(dau[0] == 2, f"2 inscrições agregadas (obtido: {dau[0]})")
            conferir(abs(dau[1] - 1_599_380.85) < 0.01,
                     "valor brasileiro convertido: 1.250.480,75 + 348.900,10",
                     f"obtido {dau[1]}")
            conferir(dau[2] is True, "indicador de ajuizamento propagado")

            fora_uf = uma("SELECT count(*) FROM divida_ativa_uniao "
                          "WHERE uf_devedor = 'SP'")[0]
            conferir(fora_uf == 0, "dívida ativa de outra UF é filtrada")

            # Dívida ativa ESTADUAL (MG) — janela de 1 ano
            total_dam = uma("SELECT count(*) FROM divida_ativa_mg")[0]
            conferir(total_dam == 2,
                     f"2 de 3 linhas carregadas (a antiga fica fora da janela de 1 ano, "
                     f"obtido: {total_dam})")

            dam_vale_verde = uma("SELECT tem_divida_ativa_mg, dam_valor_total, "
                                 "dam_especies FROM leads WHERE cnpj_basico = '10000006'")
            conferir(dam_vale_verde[0] is True,
                     "dívida estadual recente vira apontamento para Vale Verde")
            conferir(abs(dam_vale_verde[1] - 45230.80) < 0.01,
                     "valor da dívida estadual propagado", f"obtido {dam_vale_verde[1]}")
            conferir(dam_vale_verde[2] == "ICMS", "espécie do tributo estadual preservada")

            dam_metal = uma("SELECT tem_divida_ativa_mg FROM leads "
                            "WHERE cnpj_basico = '10000001'")[0]
            conferir(dam_metal is False,
                     "dívida estadual ANTIGA da metalúrgica não conta — fora da janela")

            orfao_carregado = uma("SELECT count(*) FROM divida_ativa_mg "
                                  "WHERE cnpj_basico = '99999999'")[0]
            conferir(orfao_carregado == 1,
                     "CNPJ estadual sem empresa correspondente carrega sem erro")
            orfao_vira_lead = uma("SELECT count(*) FROM leads "
                                  "WHERE cnpj_basico = '99999999'")[0]
            conferir(orfao_vira_lead == 0,
                     "mas não vira lead, por não estar no universo (LEFT JOIN correto)")

            # Situação cadastral
            conferir(uma("SELECT inapta FROM leads WHERE cnpj_basico = '10000004'")[0] is True,
                     "situação 04 marca inapta")
            conferir(uma("SELECT suspensa FROM leads WHERE cnpj_basico = '10000007'")[0] is True,
                     "situação 03 marca suspensa")
            conferir(uma("SELECT situacao_irregular FROM leads "
                         "WHERE cnpj_basico = '10000001'")[0] is False,
                     "situação 02 não é irregular")

            # Sanção
            conferir(uma("SELECT tem_sancao FROM leads WHERE cnpj_basico = '10000007'")[0] is True,
                     "sanção do CEIS é vinculada pelo CNPJ básico")

            # Ex-Simples
            conferir(uma("SELECT saiu_do_simples FROM leads "
                         "WHERE cnpj_basico = '10000006'")[0] is True,
                     "exclusão do Simples é detectada")
            conferir(uma("SELECT data_exclusao_simples FROM leads "
                         "WHERE cnpj_basico = '10000006'")[0] == "20240131",
                     "data de exclusão preservada")

            # Prognóstico de CND
            prog_metal = uma("SELECT cnd_federal_prognostico FROM leads "
                             "WHERE cnpj_basico = '10000001'")[0]
            conferir("IMPROVÁVEL" in prog_metal,
                     "dívida ativa → CND negativa improvável", prog_metal)

            prog_serra = uma("SELECT cnd_federal_prognostico FROM leads "
                             "WHERE cnpj_basico = '10000004'")[0]
            conferir("dívida ativa + inapta" in prog_serra,
                     "dívida + inapta → prognóstico combinado", prog_serra)

            prog_banco = uma("SELECT cnd_federal_prognostico FROM leads "
                             "WHERE cnpj_basico = '10000003'")[0]
            conferir(prog_banco == "SEM APONTAMENTO NAS FONTES ABERTAS",
                     "sem apontamento é rotulado como tal", prog_banco)

            # Prognóstico de CDT estadual (MG) — mesma lógica, jurisdição diferente.
            prog_mg_vale_verde = uma("SELECT cdt_estadual_mg_prognostico FROM leads "
                                     "WHERE cnpj_basico = '10000006'")[0]
            conferir("dívida ativa estadual (MG)" in prog_mg_vale_verde,
                     "dívida estadual → prognóstico de CDT-MG negativo", prog_mg_vale_verde)

            prog_mg_metal = uma("SELECT cdt_estadual_mg_prognostico FROM leads "
                                "WHERE cnpj_basico = '10000001'")[0]
            conferir(prog_mg_metal == "SEM APONTAMENTO NAS FONTES ABERTAS",
                     "sem dívida estadual (a antiga não conta) → CDT-MG sem apontamento",
                     prog_mg_metal)

            # Score: metalúrgica = dívida União (30) + porte Demais (10) = 40.
            score_metal = uma("SELECT score FROM leads WHERE cnpj_basico = '10000001'")[0]
            conferir(score_metal == 40, f"score da metalúrgica = 40 (obtido: {score_metal})")

            # Serra Azul: dívida 30 + irregular 25 + saiu do Simples 20 + Demais 10 = 85.
            score_serra = uma("SELECT score FROM leads WHERE cnpj_basico = '10000004'")[0]
            conferir(score_serra == 85, f"score da Serra Azul = 85 (obtido: {score_serra})")

            score_banco = uma("SELECT score FROM leads WHERE cnpj_basico = '10000003'")[0]
            conferir(score_banco == 10, f"score do banco = 10 (obtido: {score_banco})")

            # Vale Verde: saiu do Simples (20) + dívida estadual MG (25) = 45.
            score_vale_verde = uma("SELECT score FROM leads WHERE cnpj_basico = '10000006'")[0]
            conferir(score_vale_verde == 45,
                     f"score da Vale Verde = 45 (obtido: {score_vale_verde})")

            teto = uma("SELECT max(score) FROM leads")[0]
            conferir(teto <= 100, "score nunca passa de 100")

            # Enriquecimento por join
            conferir(uma("SELECT municipio FROM leads WHERE cnpj_basico = '10000001'")[0]
                     == "BELO HORIZONTE", "código de município resolvido para nome")
            conferir(uma("SELECT cnae_descricao FROM leads WHERE cnpj_basico = '10000001'")[0]
                     == "Fabricação de estruturas metálicas", "CNAE resolvido")
            conferir(uma("SELECT cnpj_formatado FROM leads WHERE cnpj_basico = '10000001'")[0]
                     == "10.000.001/0001-55", "CNPJ formatado")
            conferir(abs(uma("SELECT capital_social FROM leads "
                             "WHERE cnpj_basico = '10000001'")[0] - 5_000_000.0) < 0.01,
                     "capital social com vírgula decimal convertido")

            apont = uma("SELECT apontamentos_texto FROM leads "
                        "WHERE cnpj_basico = '10000004'")[0]
            conferir("Dívida ativa" in apont and "INAPTO" in apont
                     and "Excluída do Simples" in apont,
                     "texto de apontamentos concatena os três sinais", apont)

            apont_vale_verde = uma("SELECT apontamentos_texto FROM leads "
                                   "WHERE cnpj_basico = '10000006'")[0]
            conferir("Dívida ativa estadual MG" in apont_vale_verde
                     and "Excluída do Simples" in apont_vale_verde,
                     "apontamento estadual entra na lista textual", apont_vale_verde)

        print("\n5. Transformando (filtro de apontamento ON — comportamento padrão)")
        object.__setattr__(cfg, "somente_com_apontamento", True)
        transform.executar(cfg)

        with db.conectar(cfg, somente_leitura=True) as con:
            def uma(sql: str, *params):
                return con.execute(sql, list(params)).fetchone()

            total_filtrado = uma("SELECT count(*) FROM leads")[0]
            conferir(total_filtrado == 4,
                     f"filtro liga por padrão: 4 leads com apontamento (obtido: {total_filtrado})")

            banco_existe = uma("SELECT count(*) FROM leads "
                               "WHERE cnpj_basico = '10000003'")[0] > 0
            conferir(not banco_existe,
                     "banco sem apontamento é descartado quando o filtro está ligado")

            for basico in ("10000001", "10000004", "10000006", "10000007"):
                conferir(
                    uma("SELECT count(*) FROM leads WHERE cnpj_basico = ?", basico)[0] == 1,
                    f"lead {basico} com apontamento permanece após o filtro",
                )

            # A tabela auxiliar `apontamentos` continua completa (inclui o banco)
            # mesmo com o filtro ligado — só `leads`, a saída final, é cortada.
            total_apontamentos = uma("SELECT count(*) FROM apontamentos")[0]
            conferir(total_apontamentos == 5,
                     "tabela `apontamentos` continua com o universo inteiro",
                     f"obtido: {total_apontamentos}")

        print("\n6. Exportando planilha (com o filtro ligado)")
        caminho = export.executar(cfg)
        conferir(caminho.exists() and caminho.stat().st_size > 4000,
                 "planilha Excel gerada")

        from openpyxl import load_workbook
        linhas_planilha = load_workbook(caminho).active.max_row - 1  # -1 do cabeçalho
        conferir(linhas_planilha == 4,
                 f"planilha exportada já reflete o filtro: 4 linhas (obtido: {linhas_planilha})")

        print("\n7. Detecção de layout divergente")
        ruim = trabalho / "csv" / "RUIM.EMPRECSV"
        ruim.write_text("1;2;3\n4;5;6\n", encoding="latin-1")
        try:
            db.conferir_largura(ruim, schemas.EMPRESAS)
            conferir(False, "CSV com número errado de colunas é rejeitado")
        except db.LayoutInesperadoError:
            conferir(True, "CSV com número errado de colunas é rejeitado")

        print("\n8. Conversão de data serial (Excel → ISO)")
        conferir(mg_divida._serial_para_iso(45658) == "2025-01-01",
                 "serial Excel converte para ISO",
                 f"obtido {mg_divida._serial_para_iso(45658)}")
        conferir(mg_divida._serial_para_iso("não é número") is None,
                 "valor não numérico retorna None em vez de lançar exceção")
        conferir(mg_divida._serial_para_iso(None) is None,
                 "None retorna None em vez de lançar exceção")

    finally:
        shutil.rmtree(trabalho, ignore_errors=True)

    print()
    if falhas:
        print(f"{len(falhas)} verificação(ões) falharam:")
        for f in falhas:
            print(f"  - {f}")
        return 1

    print("Todas as verificações passaram.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
